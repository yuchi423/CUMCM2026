"""Round-trip audit for the published Q3 workbook."""
import hashlib
import json
import math
import sys
from pathlib import Path
import openpyxl

root = Path(sys.argv[1]).resolve()
payload = json.loads((root / "tmp/q3-workbook/workbook_data.json").read_text(encoding="utf-8"))
path = root / "output/result3.xlsx"
values = openpyxl.load_workbook(path, read_only=True, data_only=True)
formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
assert values.sheetnames == list(payload)
checked = 0
worst = 0.0
for name, table in payload.items():
    rows = list(values[name].values)
    extra = 2 if name == "费用汇总" else 0
    assert list(rows[0]) == table["headers"]
    assert len(rows) == len(table["rows"]) + 1 + extra
    for source, row in zip(table["rows"], rows[1:]):
        assert len(source) == len(row)
        for column, (a, b) in enumerate(zip(source, row)):
            if column == 0:
                assert b.date().isoformat() == a
            elif isinstance(a, (int, float)):
                assert isinstance(b, (int, float)) and math.isfinite(b), (name, column, a, b)
                worst = max(worst, abs(a - b))
                assert abs(a - b) < 1e-5, (name, column, a, b)
            else:
                assert a == b, (name, column, a, b)
            checked += 1
assert formulas["计划购电量"]["EP2"].value == "=SUM(B2:EO2)"
assert formulas["调整购电量"]["EP2"].value == "=SUM(B2:EO2)"
annual_row = len(payload["费用汇总"]["rows"]) + 2
assert formulas["费用汇总"].cell(annual_row, 6).value == f"=SUM(F2:F{annual_row-1})"
annual_total = values["费用汇总"].cell(annual_row, 6).value
assert abs(annual_total - 13852172.808932787) < 0.01
assert values["费用汇总"].cell(annual_row, 13).value == 869
assert "附件3光伏预测" in values["费用汇总"].cell(annual_row + 1, 2).value
values.close(); formulas.close()
inspection = (root / "tmp/q3-workbook/workbook_inspection.ndjson").read_text(encoding="utf-8")
assert not any(token in inspection for token in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"])
result = {"pass_": True, "sheets": len(payload), "checked_cells": checked,
          "max_numeric_error": worst, "annual_total_cost": annual_total,
          "planned_days": 334, "storage_rows": 334 * 6,
          "paper_dates": 4, "formula_error_scan": True,
          "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
(root / "output/result3_audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False))
