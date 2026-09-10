"""Read-only Q1 input audit. Does not clean data, simulate or optimize."""
import argparse
import csv
import hashlib
import importlib.metadata
import json
import math
import platform
import subprocess
from datetime import datetime, time, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def minute_of_day(value):
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"Non-minute timestamp: {value}")
        return 60 * value.hour + value.minute
    if str(value).strip() in {"0:00+1", "00:00+1", "24:00"}:
        return 1440
    raise ValueError(f"Unsupported time label: {value!r}")


def clock_label(minutes):
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    if not output.is_relative_to((ROOT / "experiments").resolve()):
        raise ValueError("Audit output must be inside experiments/")
    if output.exists():
        raise FileExistsError("Use a unique audit output path; do not overwrite evidence")
    with (ROOT / "data/MANIFEST.csv").open(encoding="utf-8-sig", newline="") as handle:
        manifest = {row["path_or_uri"]: row["sha256"] for row in csv.DictReader(handle)}
    paths = ["problem/C题.pdf", "data/raw/附件1.xlsx", "data/templates/result1.xlsx"]
    hashes = {path: {"actual": digest(ROOT / path), "expected": manifest[path]} for path in paths}
    for record in hashes.values():
        record["match"] = record["actual"] == record["expected"]
    if not all(record["match"] for record in hashes.values()):
        raise ValueError("Q1 input hash mismatch")

    book = openpyxl.load_workbook(ROOT / paths[1], read_only=True, data_only=True)
    sheet = book["Sheet1"]
    header = [cell.value for cell in sheet[1]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(rows) != 144 or len(header) != 4:
        raise ValueError("Expected 144 input records and 4 columns")
    minutes = [minute_of_day(row[0]) for row in rows]
    if minutes != list(range(10, 1441, 10)):
        raise ValueError("Time axis is not exactly 00:10 through 24:00 at 10-minute spacing")
    columns = {}
    for index, name in enumerate(header[1:], start=1):
        values = [row[index] for row in rows]
        invalid = [i + 2 for i, val in enumerate(values)
                   if isinstance(val, bool) or not isinstance(val, (float, int)) or not math.isfinite(val)]
        if invalid:
            raise ValueError(f"Missing/nonfinite/nonnumeric {name} in rows {invalid}")
        columns[str(name)] = {
            "minimum": min(values), "maximum": max(values),
            "negative_count": sum(val < 0 for val in values),
            "zero_count": sum(val == 0 for val in values),
            "missing_or_nonfinite_count": 0,
        }
    book.close()

    template = openpyxl.load_workbook(ROOT / paths[2], read_only=True, data_only=False)
    plan = template["计划购电量"]
    if plan.max_row != 145:
        raise ValueError("Q1 template must have 144 plan rows")
    mapping = []
    for index, (source, end) in enumerate(zip(rows, minutes), start=2):
        mapping.append({
            "source_cell": f"Sheet1!A{index}",
            "source_time": str(source[0]),
            "proposed_interval_if_end_timestamp": f"{clock_label(end - 10)}-{clock_label(end)}",
            "template_cell": f"计划购电量!A{index}",
            "template_original_label": plan.cell(index, 1).value,
            "template_value_cell": f"计划购电量!B{index}",
            "mapping_status": "PROPOSED_NOT_APPROVED",
        })
    template_info = {
        "sheets": template.sheetnames,
        "plan_output_nonblank_count": sum(plan.cell(i, 2).value is not None for i in range(2, 146)),
        "storage_sheet_rows": list(template["充放电量"].iter_rows(values_only=True)),
        "interval_mapping_proposal": mapping,
    }
    template.close()
    versions = {"python": platform.python_version(), "openpyxl": openpyxl.__version__}
    for package in ["numpy", "scipy"]:
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "NOT_INSTALLED"
    report = {
        "scope": "Q1 raw input and template audit only; no numerical optimization or PoC",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "code_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "versions": versions,
        "input_hashes": hashes,
        "data": {
            "sheet": "Sheet1", "header": header, "records": len(rows),
            "time_sequence_check": "PASS", "numeric_check": "PASS", "columns": columns,
            "pv_exceeds_load_count": sum(row[3] > row[2] for row in rows),
            "power_sample_semantics": "UNRESOLVED: average power versus instantaneous sample",
        },
        "template": template_info,
        "poc": "NOT_EXECUTED",
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output.relative_to(ROOT)), "data": report["data"],
                      "versions": versions, "template_first": mapping[0], "template_last": mapping[-1],
                      "poc": "NOT_EXECUTED"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    main()
