"""Read back result4-3.xlsx and reconcile every populated value to the audited payload."""
from __future__ import annotations

import hashlib
import json
import math
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def date_text(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value.isoformat() if hasattr(value, "isoformat") else value


def main(run_arg: str) -> None:
    run = (ROOT / run_arg).resolve()
    payload = json.loads((run / "workbook_data.json").read_text(encoding="utf-8"))
    target = ROOT / payload["output"]
    template = ROOT / payload["template"]
    original = openpyxl.load_workbook(template, read_only=False, data_only=False)
    book = openpyxl.load_workbook(target, read_only=False, data_only=False)
    errors = []
    if book.sheetnames != original.sheetnames:
        errors.append("sheet names/order changed")
    checked = 0
    max_error = 0.0
    for name, expected_rows in payload["sheets"].items():
        ws, source = book[name], original[name]
        width = len(expected_rows[0])
        actual_header = [ws.cell(1, col).value for col in range(1, width + 1)]
        source_header = [source.cell(1, col).value for col in range(1, width + 1)]
        if actual_header != source_header:
            errors.append(f"{name} header changed")
        for row_no, expected in enumerate(expected_rows, 2):
            for col_no, wanted in enumerate(expected, 1):
                actual = ws.cell(row_no, col_no).value
                if col_no == 1:
                    actual = date_text(actual)
                if isinstance(wanted, (int, float)):
                    if not isinstance(actual, (int, float)) or not math.isfinite(actual):
                        errors.append(f"{name}!{row_no},{col_no} not numeric")
                    else:
                        delta = abs(float(actual) - float(wanted)); max_error = max(max_error, delta)
                        if delta > 1e-5:
                            errors.append(f"{name}!{row_no},{col_no} delta={delta}")
                elif actual != wanted:
                    errors.append(f"{name}!{row_no},{col_no}: {actual!r} != {wanted!r}")
                checked += 1
    book.close(); original.close()
    result = {
        "pass": not errors, "checked_cells": checked, "max_numeric_error": max_error,
        "headers_equal_official_template": not any("header" in error for error in errors),
        "sheet_names_and_order_equal_template": "sheet names/order changed" not in errors,
        "planned_days": payload["counts"]["days"],
        "storage_rows": payload["counts"]["storage_rows"],
        "emergency_rows": payload["counts"]["emergency_rows"],
        "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
        "output_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "errors": errors[:20]
    }
    (run / "workbook_audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if errors:
        raise AssertionError(result)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])
