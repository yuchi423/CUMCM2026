"""Inspect workbook dimensions and labels without modifying raw inputs."""
import importlib.metadata
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def main():
    reports = []
    files = sorted((ROOT / 'data/raw').glob('*.xlsx')) + sorted((ROOT / 'data/templates').glob('*.xlsx'))
    for path in files:
        if path.name.startswith('~$'):
            continue
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows)
                first = next(rows)
                last = first
                for row in rows:
                    last = row
                reports.append({
                    'file': path.relative_to(ROOT).as_posix(), 'sheet': sheet.title,
                    'rows_including_header': sheet.max_row, 'columns': sheet.max_column,
                    'header_first_5': list(header[:5]), 'header_last_4': list(header[-4:]),
                    'first_data_row_first_4': list(first[:4]),
                    'last_data_row_first_2': list(last[:2]),
                })
        finally:
            book.close()
    result = {
        'scope': 'Workbook structure and boundary labels only; not a data-quality or model validation.',
        'openpyxl_version': importlib.metadata.version('openpyxl'), 'sheets': reports,
    }
    target = ROOT / 'data/structure-audit.json'
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
    print(f'Workbook structure: {len(files)} workbooks, {len(reports)} sheets; {target.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
