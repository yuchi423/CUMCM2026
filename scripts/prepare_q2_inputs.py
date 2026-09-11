"""Read Q2 source workbooks without mutation; produce a bounded JSON input snapshot."""
import argparse
import csv
import hashlib
import json
import math
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def minute(value):
    if hasattr(value, "hour"):
        return value.hour * 60 + value.minute
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::00)?(\+1)?", str(value).strip())
    if not match:
        raise ValueError(f"Unrecognized timestamp: {value!r}")
    return int(match[1]) * 60 + int(match[2]) + (1440 if match[3] else 0)


def day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(str(value).split()[0].replace("/", "-")).isoformat()


def read_sources(config):
    with (ROOT / "data/MANIFEST.csv").open(encoding="utf-8-sig", newline="") as file:
        expected = {r["path_or_uri"]: r["sha256"] for r in csv.DictReader(file)}
    paths = ["problem/C题.pdf", "data/raw/附件1.xlsx", "data/raw/附件2.xlsx", "data/templates/result2.xlsx"]
    hashes = {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in paths}
    assert all(hashes[p] == expected[p] for p in paths), "Input hash mismatch"
    book = openpyxl.load_workbook(ROOT / paths[1], read_only=True, data_only=True)
    raw = list(book["Sheet1"].iter_rows(min_row=2, values_only=True))
    assert len(raw) == 144 and [minute(r[0]) for r in raw] == list(range(10, 1441, 10))
    prices = [float(r[1]) for r in raw]
    assert all(math.isfinite(p) and p > 0 for p in prices)
    book.close()
    book = openpyxl.load_workbook(ROOT / paths[2], read_only=True, data_only=True)
    curves = []
    for name in ["小区负载", "光伏发电实际功率"]:
        rows = iter(book[name].iter_rows(values_only=True))
        header = next(rows)
        assert len(header) == 145 and [minute(x) for x in header[1:]] == list(range(10, 1441, 10))
        data = {}
        for row in rows:
            label = day(row[0])
            if not config["warmup_start"] <= label <= config["evaluation_end"]:
                continue
            assert label not in data, f"Duplicate date {label}"
            values = row[1:]
            assert len(values) == 144
            assert all(isinstance(x, (int, float)) and not isinstance(x, bool)
                       and math.isfinite(x) and x >= 0 for x in values), (name, label)
            data[label] = [float(x) * config["step_hours"] for x in values]
        curves.append(data)
    book.close()
    dates = sorted(curves[0])
    assert dates == sorted(curves[1])
    assert dates[0] == config["warmup_start"] and dates[-1] == config["evaluation_end"]
    assert all((date.fromisoformat(b) - date.fromisoformat(a)).days == 1 for a, b in zip(dates, dates[1:]))
    template = openpyxl.load_workbook(ROOT / paths[3], read_only=True, data_only=False)
    plan = template["计划购电量"]
    mapping = [{"source_end_minute": (t + 1) * 10,
                "physical_start_minute": t * 10,
                "physical_end_minute": (t + 1) * 10,
                "original_template_label": str(plan.cell(1, t + 2).value)} for t in range(144)]
    template.close()
    return {"dates": dates, "prices": prices,
            "load": [curves[0][d] for d in dates], "pv": [curves[1][d] for d in dates],
            "input_hashes": hashes, "template_mapping": mapping,
            "audit": {"checked_days": len(dates), "slots_per_day": 144, "time_and_numeric_check": "PASS",
                      "scope": "Only the bounded warmup/PoC days; later dates not numerically audited"}}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--config", type=Path, default=ROOT / "configs/q2_poc.json")
    args = parser.parse_args()
    target = args.output.resolve()
    assert target.is_relative_to((ROOT / "experiments").resolve())
    if target.exists():
        raise FileExistsError("Choose a unique run directory")
    config = json.loads(args.config.read_text(encoding="utf-8"))
    data = read_sources(config)
    target.mkdir(parents=True)
    (target / "inputs.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(data["audit"]))
